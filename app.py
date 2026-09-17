from __future__ import annotations

import logging
import os
from pathlib import Path

from dotenv import load_dotenv
from docx import Document
from openai import OpenAI
from openpyxl import load_workbook
from pypdf import PdfReader
from telegram import Update
from telegram.ext import Application, CommandHandler, ContextTypes, MessageHandler, filters

from memory import MemoryStore
from prompts import SYSTEM_INSTRUCTIONS


load_dotenv()
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")

TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
MODEL = os.environ.get("OPENAI_MODEL", "gpt-5-mini").strip()
DATA_DIR = Path(os.environ.get("DATA_DIR", "./data")).resolve()
FILES_DIR = DATA_DIR / "files"
FILES_DIR.mkdir(parents=True, exist_ok=True)
ALLOWED_IDS = {
    int(x.strip())
    for x in os.environ.get("ALLOWED_TELEGRAM_USER_IDS", "").split(",")
    if x.strip().isdigit()
}

store = MemoryStore(DATA_DIR / "agent.db")
client = OpenAI() if os.environ.get("OPENAI_API_KEY") else None


def allowed(update: Update) -> bool:
    user = update.effective_user
    return bool(user and (not ALLOWED_IDS or user.id in ALLOWED_IDS))


async def reject(update: Update):
    if update.effective_message:
        await update.effective_message.reply_text("هذا البوت خاص وغير مصرح لهذا الحساب باستخدامه.")


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    if suffix == ".docx":
        doc = Document(path)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"[Sheet: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
        return "\n".join(parts)
    if suffix in {".txt", ".md", ".csv"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError("النوع المدعوم حاليًا: PDF، DOCX، XLSX، TXT، MD، CSV")


def build_prompt(user_id: int, request: str) -> str:
    memories = store.list_memories(user_id, 30)
    context = store.search_context(user_id, request)
    recent = store.recent_messages(user_id, 10)
    memory_text = "\n".join(f"- [{m['id']}] {m['content']}" for m in memories)
    context_text = "\n\n".join(
        f"المصدر: {row['label']}\n{row['text']}" for row in context
    )
    history_text = "\n".join(f"{r['role']}: {r['content']}" for r in recent)
    return f"""
التعليمات الدائمة المحفوظة:
{memory_text or 'لا توجد تعليمات محفوظة.'}

السياق المسترجع من الذاكرة والملفات:
{context_text or 'لا يوجد سياق مطابق.'}

آخر المحادثة:
{history_text or 'لا يوجد.'}

طلب المستخدم الحالي:
{request}
""".strip()


def ask_ai(user_id: int, request: str) -> str:
    if not client:
        return "تم الاستقبال، لكن مفتاح نموذج الذكاء الاصطناعي غير مضاف بعد. أضف OPENAI_API_KEY ثم أعد التشغيل."
    response = client.responses.create(
        model=MODEL,
        instructions=SYSTEM_INSTRUCTIONS,
        input=build_prompt(user_id, request),
    )
    return response.output_text.strip()


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not allowed(update):
        return await reject(update)
    await update.message.reply_text(
        "مرحبًا، أنا وكيل أعمالك. أرسل سؤالًا أو ملفًا، أو استخدم:\n"
        "/myid رقم حسابك\n"
        "/remember نص لحفظ تعليمات دائمة\n"
        "/memories عرض الذاكرة\n"
        "/forget رقم حذف معلومة\n"
        "/help المساعدة"
    )


async def myid(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(f"رقم حساب تيليجرام: {update.effective_user.id}")


async def remember(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not allowed(update):
        return await reject(update)
    text = " ".join(context.args).strip()
    if not text:
        return await update.message.reply_text("اكتب: /remember ثم التعليمات التي تريد حفظها.")
    memory_id = store.add_memory(update.effective_user.id, text)
    await update.message.reply_text(f"تم حفظها في الذاكرة الدائمة برقم {memory_id}.")


async def memories(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not allowed(update):
        return await reject(update)
    rows = store.list_memories(update.effective_user.id)
    if not rows:
        return await update.message.reply_text("الذاكرة الدائمة فارغة.")
    text = "\n".join(f"{r['id']}. {r['content']}" for r in rows)
    await update.message.reply_text(text[:4000])


async def forget(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not allowed(update):
        return await reject(update)
    if not context.args or not context.args[0].isdigit():
        return await update.message.reply_text("اكتب: /forget ثم رقم المعلومة.")
    deleted = store.delete_memory(update.effective_user.id, int(context.args[0]))
    await update.message.reply_text("تم الحذف." if deleted else "لم أجد هذه المعلومة.")


async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not allowed(update):
        return await reject(update)
    user_id = update.effective_user.id
    text = update.message.text.strip()
    store.add_message(user_id, "المستخدم", text)
    await update.message.chat.send_action("typing")
    answer = ask_ai(user_id, text)
    store.add_message(user_id, "الوكيل", answer)
    await update.message.reply_text(answer[:4000])


async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not allowed(update):
        return await reject(update)
    doc = update.message.document
    safe_name = Path(doc.file_name or f"file-{doc.file_unique_id}").name
    user_dir = FILES_DIR / str(update.effective_user.id)
    user_dir.mkdir(parents=True, exist_ok=True)
    target = user_dir / f"{doc.file_unique_id}-{safe_name}"
    tg_file = await doc.get_file()
    await tg_file.download_to_drive(custom_path=target)
    try:
        text = extract_text(target)
    except Exception as exc:
        target.unlink(missing_ok=True)
        return await update.message.reply_text(f"تعذر قراءة الملف: {exc}")
    if not text.strip():
        return await update.message.reply_text("تم استلام الملف، لكن لم أجد نصًا قابلاً للاستخراج. قد يكون PDF ممسوحًا ضوئيًا ويحتاج OCR.")
    doc_id = store.add_document(update.effective_user.id, safe_name, str(target), text[:2_000_000])
    caption = (update.message.caption or "حلل الملف وقدّم ملخصًا تنفيذيًا وأهم المهام والمخاطر.").strip()
    answer = ask_ai(update.effective_user.id, f"الملف المرفوع: {safe_name}، رقم الوثيقة {doc_id}. {caption}")
    await update.message.reply_text(f"تم حفظ الملف برقم {doc_id}.\n\n{answer}"[:4000])


def main():
    if not TOKEN:
        raise RuntimeError("TELEGRAM_BOT_TOKEN غير موجود. انسخ .env.example إلى .env وأضف الرمز.")
    app = Application.builder().token(TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("help", start))
    app.add_handler(CommandHandler("myid", myid))
    app.add_handler(CommandHandler("remember", remember))
    app.add_handler(CommandHandler("memories", memories))
    app.add_handler(CommandHandler("forget", forget))
    app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text))
    app.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
